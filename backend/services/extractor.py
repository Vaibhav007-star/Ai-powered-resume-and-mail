import re
import json
import logging
from pathlib import Path
from typing import Dict, Any, Optional
import pypdf
import docx

logger = logging.getLogger(__name__)

def extract_text_from_file(file_path: Path) -> str:
    """Extract plain text from .pdf, .docx, or .txt file."""
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")
    
    suffix = file_path.suffix.lower()
    
    if suffix == ".pdf":
        return _extract_from_pdf(file_path)
    elif suffix in (".docx", ".doc"):
        return _extract_from_docx(file_path)
    elif suffix in (".txt", ".md"):
        return _extract_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file format '{suffix}'. Supported formats: .pdf, .docx, .txt")

def _extract_from_pdf(file_path: Path) -> str:
    extracted_text = []
    with open(file_path, "rb") as f:
        reader = pypdf.PdfReader(f)
        for page_num, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                extracted_text.append(text.strip())
    return "\n\n".join(extracted_text)

def _extract_from_docx(file_path: Path) -> str:
    doc = docx.Document(file_path)
    full_text = []
    
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())
            
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                full_text.append(" | ".join(row_text))
                
    return "\n".join(full_text)

def _extract_from_txt(file_path: Path) -> str:
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1") as f:
            return f.read()

RE_EMAIL = re.compile(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+")
RE_PHONE = re.compile(r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}")
DEGREE_PATTERNS = [
    (d, re.compile(r"\b" + re.escape(d) + r"\b", re.IGNORECASE))
    for d in ["Bachelor of Science", "Bachelor of Engineering", "Bachelor of Technology", "B.Tech", "B.E.", "B.S.", "Master of Science", "Master of Technology", "M.Tech", "M.S.", "MBA", "Ph.D."]
]

def parse_profile_heuristically(raw_text: str) -> Dict[str, Any]:
    """Fallback heuristic extractor when LLM key is not provided."""
    lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    
    email_match = RE_EMAIL.search(raw_text)
    email = email_match.group(0) if email_match else ""
    
    phone_match = RE_PHONE.search(raw_text)
    phone = phone_match.group(0) if phone_match else ""
    
    full_name = lines[0] if lines and len(lines[0].split()) <= 4 and "@" not in lines[0] else ""
    
    degree = ""
    for d_name, d_regex in DEGREE_PATTERNS:
        if d_regex.search(raw_text):
            degree = d_name
            break
            
    return {
        "full_name": full_name,
        "email": email,
        "phone": phone,
        "degree": degree,
        "graduation_year": "",
        "skills": [],
        "experience": [],
        "projects": [],
        "certifications": [],
        "preferred_roles": [],
        "preferred_locations": [],
        "work_mode": "Any",
        "experience_level": "Mid",
        "other_preferences": "",
        "raw_resume_text": raw_text
    }

def extract_structured_profile(raw_text: str) -> Dict[str, Any]:
    """Extract structured candidate profile from raw resume text."""
    if not raw_text.strip():
        return parse_profile_heuristically("")

    from backend.core.llm import is_llm_configured, generate_json_response

    if not is_llm_configured():
        logger.info("No AI API key configured. Using heuristic parsing fallback.")
        return parse_profile_heuristically(raw_text)

    prompt = f"""You are a precise, truth-preserving resume parser.

CRITICAL ACCURACY RULE:
- NEVER INVENT, ASSUME, OR EXTRAPOLATE INFORMATION.
- Extract ONLY what is explicitly stated in the resume text below.
- If any detail is missing, leave it as an empty string or empty list.
- Do NOT add technologies, skills, or titles that are not directly mentioned.

Resume Text:
---
{raw_text}
---

Extract the information into the following JSON schema:
{{
  "full_name": "string (candidate full name or empty)",
  "email": "string (email address or empty)",
  "phone": "string (phone number or empty)",
  "degree": "string (e.g., B.Tech in Computer Science or empty)",
  "graduation_year": "string (e.g. 2024 or empty)",
  "skills": ["list of explicit technical and professional skills"],
  "experience": [
    {{
      "title": "job title",
      "company": "company name",
      "duration": "employment duration / dates",
      "description": "brief summary of responsibilities as stated"
    }}
  ],
  "projects": [
    {{
      "name": "project title",
      "tech": "technologies explicitly used in project",
      "description": "summary of the project as stated"
    }}
  ],
  "certifications": ["list of explicit certifications"],
  "experience_level": "Entry | Mid | Senior | Lead (inferred strictly from years of experience stated, default 'Mid')"
}}
"""
    try:
        extracted_data = generate_json_response(prompt=prompt, temperature=0.0)
        if extracted_data and isinstance(extracted_data, dict):
            extracted_data["raw_resume_text"] = raw_text
            return extracted_data
    except Exception as e:
        logger.error(f"Error during AI resume extraction: {e}. Falling back to heuristics.")

    fallback = parse_profile_heuristically(raw_text)
    return fallback
