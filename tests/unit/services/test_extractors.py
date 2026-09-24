import os
import sys
from pathlib import Path
import pytest
from pypdf import PdfWriter
import docx

BASE_DIR = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(BASE_DIR))

from backend.services.extractor import extract_text_from_file

def test_extract_from_txt(tmp_path):
    txt_file = tmp_path / "resume.txt"
    txt_file.write_text("Alex Smith\nPython Developer\nSkills: FastAPI, SQL", encoding="utf-8")
    extracted = extract_text_from_file(txt_file)
    assert "Alex Smith" in extracted
    assert "FastAPI" in extracted

def test_extract_from_docx(tmp_path):
    docx_file = tmp_path / "resume.docx"
    doc = docx.Document()
    doc.add_heading("Maria Garcia", level=1)
    doc.add_paragraph("maria.garcia@example.com")
    doc.add_paragraph("Experience: 4 years as Backend Engineer")
    doc.save(docx_file)

    extracted = extract_text_from_file(docx_file)
    assert "Maria Garcia" in extracted
    assert "maria.garcia@example.com" in extracted
    assert "Backend Engineer" in extracted

def test_extract_from_pdf(tmp_path):
    pdf_file = tmp_path / "resume.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=72 * 8.5, height=72 * 11)
    # Write blank PDF and verify it doesn't crash on empty or valid structure
    with open(pdf_file, "wb") as f:
        writer.write(f)

    extracted = extract_text_from_file(pdf_file)
    assert isinstance(extracted, str)
